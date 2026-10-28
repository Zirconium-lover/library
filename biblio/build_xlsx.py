"""Сборка biblio/catalog.xlsx из biblio/catalog.csv (листы «Реестр», «Темы», «Параметры», «Сводка»)."""
import csv, os, re
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

B = os.path.dirname(os.path.abspath(__file__))

TOPICS = {
    "1.1": "Вязкое разрушение в континууме: CDM, Rice–Tracey, Johnson–Cook, Bai–Wierzbicki, локусы εf(η, θ)",
    "1.2": "Пористая пластичность: Gurson, GTN, Thomason, Nahshon–Hutchinson, слияние пор, нелокальный GTN",
    "1.3": "Регуляризация и масштаб: crack band, нелокальная, градиентная, микроморфная, фазовое поле, ℓ, size effect",
    "1.4": "Когезионные зоны: законы, смешанный режим, штрафная жёсткость, вязкая регуляризация, большие повороты",
    "1.5": "Хрупкое разрушение включений: главное напряжение, crack band, Вейбулл, растрескивание частиц",
    "1.6": "Численные алгоритмы: IMPL-EX, arc-length, line search, staggered/monolithic, XFEM, удаление элементов",
    "1.7": "Микромеханика и RVE с включениями и границей: отслоение, гомогенизация разупрочнения, размер RVE",
    "1.8": "Верификация и бенчмарки: Sandia Fracture Challenge, эталоны для когезии и нелокальных моделей",
    "2.1": "Свойства гидридов: упругость, твёрдость, K_IC, прочность, микроконсоли и микростолбы",
    "2.2": "Граница гидрид–матрица: микромеханика, работа разделения, отслоение, поры",
    "2.3": "Матрица: упругопластичность, анизотропия, упрочнение, εf(η, θ); Zry, ZIRLO, M5, Э110, Э635, Zr-2.5Nb",
    "2.4": "Механизмы разрушения гидрированных сплавов: скол гидрида, поры, in situ SEM/DIC/томография",
    "2.5": "Испытания гидрированных оболочек: RCT, burst, EDC, PLT, вязко-хрупкий переход, радиальные гидриды",
    "2.6": "Деформация превращения и остаточные напряжения вокруг гидридов",
    "2.7": "Морфология и меры связности: RHCF, Fn, HCC, анализ изображений, 3D",
    "2.8": "DHC — разрушение гидрида у вершины трещины: K_IH, критический размер",
    "3.1": "Явные гидриды в МКЭ: когезия, XFEM, фазовое поле, кристаллическая пластичность",
    "3.2": "Континуальные модели повреждения гидрированного Zr",
    "3.3": "Перколяция, критический путь, стохастические модели, машинное обучение",
    "3.4": "Механика разрушения и инженерные критерии для радиальных гидридов; модели DHC",
    "3.5": "Топливные коды: модули отказа оболочки по гидридам",
}

HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="305496")
WRAP = Alignment(wrap_text=True, vertical="top")


def sheet(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD, FILL, WRAP
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def split_value(frag, rid, material_hint):
    frag = frag.strip()
    cond = "; ".join(re.findall(r"\(([^()]*)\)", frag))
    core = re.sub(r"\([^()]*\)", "", frag).strip()
    if "=" in core:
        par, val = core.split("=", 1)
    elif ":" in core:
        par, val = core.split(":", 1)
    else:
        par, val = core, ""
    m = re.match(r"\s*([≈~<>≤≥]?\s*[-+−]?[\d.,]+(?:\s*(?:±|–|-|…)\s*[\d.,]+)?(?:\s*[×x·]\s*10\^?[-−]?\d+)?)\s*(.*)", val)
    num, unit = (m.group(1).strip(), m.group(2).strip()) if m else (val.strip(), "")
    return [par.strip(), num, unit, cond, material_hint, rid]


def main():
    with open(os.path.join(B, "catalog.csv"), encoding="utf-8", newline="") as f:
        rows = list(csv.DictReader(f))
    wb = Workbook()
    ws = wb.active
    ws.title = "Реестр"
    reg = []
    for r in rows:
        reg.append([int(r["id"]), int(r["year"]) if r["year"].isdigit() else r["year"], r["title"],
                    f'{r["direction"]} / {r["subtopics"]}', r["summary_ru"], r["authors"], r["journal"],
                    r["type"], r["doi"] or r["url"], r["access"], r["priority"], r["useful_for"], r["values"]])
    sheet(ws, ["№", "Год", "Название", "Направление/подтемы", "Краткое содержание", "Авторы", "Журнал", "Тип",
               "DOI", "Доступ", "Приоритет", "Что берём для модели", "Значения"], reg,
          [6, 7, 50, 14, 70, 30, 28, 14, 28, 8, 9, 40, 45])
    cnt = Counter(s.strip() for r in rows for s in r["subtopics"].split(";") if s.strip())
    sheet(wb.create_sheet("Темы"), ["Подтема", "Что искать", "Число статей"],
          [[k, v, cnt.get(k, 0)] for k, v in TOPICS.items()], [9, 90, 12])
    par = []
    for r in rows:
        if "2" not in r["direction"].split(";") or not r["values"].strip():
            continue
        mat = ", ".join(sorted(set(re.findall(r"(Zry-?\d|Zircaloy-?\d|ZIRLO|M5|Э110|Э635|E110|E635|Zr-2\.5Nb|Zr-1Nb|δ-?ZrH[\d.]*|γ-?ZrH|ε-?ZrH[\d.]*|δ-гидрид|ZrH[\d.x]*)", r["values"]))))
        for frag in r["values"].split(" | "):
            for sub in frag.split("; "):
                if re.search(r"\d", sub):
                    par.append(split_value(sub, int(r["id"]), mat) + [r["values_source"]])
    sheet(wb.create_sheet("Параметры"), ["Параметр", "Значение", "Единицы", "Условия", "Материал", "id", "Источник числа"],
          par, [40, 18, 18, 45, 18, 6, 12])
    s = []
    s.append(["Всего записей", len(rows)])
    for d in "123":
        s.append([f"Направление {d}", sum(d in r["direction"].split(";") for r in rows)])
    for k, v in sorted(Counter(r["doi_check"] for r in rows).items()):
        s.append([f"doi_check = {k or '—'}", v])
    for k, v in sorted(Counter(r["priority"] for r in rows).items()):
        s.append([f"Приоритет {k or '—'}", v])
    for k, v in sorted(Counter(r["access"] for r in rows).items()):
        s.append([f"Доступ {k or '—'}", v])
    for k, v in Counter(r["type"] for r in rows).most_common():
        s.append([f"Тип: {k or '—'}", v])
    s.append(["Пересечение с каталогом переориентации", sum(bool(r["overlap_reorient"]) for r in rows)])
    s.append(["PDF в library", sum(bool(r["pdf_in_library"]) for r in rows)])
    s.append(["Строк в листе «Параметры»", len(par)])
    sheet(wb.create_sheet("Сводка"), ["Показатель", "Значение"], s, [45, 12])
    wb.save(os.path.join(B, "catalog.xlsx"))
    print(f"catalog.xlsx: {len(rows)} записей, {len(par)} параметров")


if __name__ == "__main__":
    main()
